from __future__ import annotations

import csv
import io
import math
import re
import unicodedata
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import EmailStr, TypeAdapter, ValidationError
from sqlmodel import Session, col, select

from app.core.config import settings
from app.models import (
    ActivationMethod,
    AthleteImportFailure,
    AthleteImportPreview,
    AthleteImportPreviewRow,
    AthleteImportResult,
    AthleteProfile,
    CoachAssignment,
    Gym,
    GymMembership,
    GymMembershipRole,
    User,
)
from app.services.membership import activate_user, add_roles, ensure_membership
from app.services.profiles import coach_members

MAX_IMPORT_BYTES = settings.ATHLETE_IMPORT_MAX_BYTES
MAX_IMPORT_ROWS = 500
MAX_READ_ROWS = 2000

CANONICAL_COLUMNS = [
    "email",
    "full_name",
    "phone",
    "coach_email",
    "training_goal",
    "height_cm",
    "weight_kg",
    "training_days_count",
]

COLUMN_ALIASES = {
    "email": "email",
    "e_mail": "email",
    "ایمیل": "email",
    "full_name": "full_name",
    "fullname": "full_name",
    "name": "full_name",
    "نام": "full_name",
    "نام_کامل": "full_name",
    "نام_و_نام_خانوادگی": "full_name",
    "phone": "phone",
    "phone_number": "phone",
    "mobile": "phone",
    "تلفن": "phone",
    "موبایل": "phone",
    "شماره_تماس": "phone",
    "coach_email": "coach_email",
    "coach": "coach_email",
    "مربی": "coach_email",
    "ایمیل_مربی": "coach_email",
    "training_goal": "training_goal",
    "goal": "training_goal",
    "هدف": "training_goal",
    "هدف_تمرینی": "training_goal",
    "height_cm": "height_cm",
    "height": "height_cm",
    "قد": "height_cm",
    "قد_سانتی_متر": "height_cm",
    "weight_kg": "weight_kg",
    "weight": "weight_kg",
    "وزن": "weight_kg",
    "وزن_کیلوگرم": "weight_kg",
    "training_days_count": "training_days_count",
    "training_days": "training_days_count",
    "days_per_week": "training_days_count",
    "تعداد_روز_تمرین": "training_days_count",
    "روزهای_تمرین": "training_days_count",
    "تعداد_روز_در_هفته": "training_days_count",
}

DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
_EMAIL = TypeAdapter(EmailStr)

SAMPLE_ROW = {
    "email": "athlete@example.com",
    "full_name": "نام و نام خانوادگی ورزشکار",
    "phone": "09120000000",
    "coach_email": "coach@example.com",
    "training_goal": "افزایش قدرت و حجم عضلانی",
    "height_cm": "180",
    "weight_kg": "80",
    "training_days_count": "3",
}

COLUMN_GUIDE = [
    ("email", True, "ایمیل یکتای ورزشکار (ستون اجباری)", "athlete@example.com"),
    ("full_name", False, "نام و نام خانوادگی", "علی رضایی"),
    ("phone", False, "شماره تماس", "09120000000"),
    (
        "coach_email",
        False,
        "ایمیل مربی؛ اگر خالی باشد از مربی پیش‌فرض انتخابی استفاده می‌شود",
        "coach@example.com",
    ),
    ("training_goal", False, "هدف تمرینی", "قدرت و حجم"),
    ("height_cm", False, "قد به سانتی‌متر", "180"),
    ("weight_kg", False, "وزن به کیلوگرم", "80"),
    ("training_days_count", False, "تعداد روز تمرین در هفته (۱ تا ۷)", "3"),
]


@dataclass
class _RawRow:
    row_number: int
    values: dict[str, str]


@dataclass
class _PreparedRow:
    row_number: int
    email: str
    full_name: str | None
    phone: str | None
    coach_email: str | None
    coach: User | None
    training_goal: str | None
    height_cm: float | None
    weight_kg: float | None
    training_days_count: int | None
    status: Literal["valid", "skip", "error"]
    messages: list[str] = field(default_factory=list)


def _normalise_header(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _as_text(value)).lower()
    text = re.sub(r"[\s\-]+", "_", text).strip("_")
    return COLUMN_ALIASES.get(text, text)


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_number(value: str) -> float | None:
    text = value.translate(DIGITS).replace(",", "").replace("٫", ".").strip()
    if not text:
        return None
    number = float(text)
    if not math.isfinite(number):
        raise ValueError
    return number


def _read_csv(content: bytes) -> list[list[Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1256")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=400, detail="invalid_file_encoding") from exc
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[Any]] = []
    try:
        for row in csv.reader(io.StringIO(text), dialect, strict=True):
            rows.append(list(row))
            if len(rows) > MAX_READ_ROWS:
                raise HTTPException(status_code=400, detail="too_many_rows")
    except csv.Error as exc:
        raise HTTPException(status_code=400, detail="invalid_file") from exc
    return rows


def _read_xlsx(content: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
    except Exception as exc:
        raise HTTPException(status_code=400, detail="invalid_file") from exc
    sheet = workbook.worksheets[0]
    rows: list[list[Any]] = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(list(row))
        if len(rows) > MAX_READ_ROWS:
            workbook.close()
            raise HTTPException(status_code=400, detail="too_many_rows")
    workbook.close()
    return rows


def _parse_rows(filename: str, content: bytes) -> list[_RawRow]:
    if not content:
        raise HTTPException(status_code=400, detail="empty_file")
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=400, detail="import_file_too_large")

    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        table = _read_csv(content)
    elif suffix == ".xlsx":
        table = _read_xlsx(content)
    else:
        raise HTTPException(status_code=400, detail="invalid_file_type")

    if not table:
        raise HTTPException(status_code=400, detail="empty_file")

    headers = [_normalise_header(value) for value in table[0]]
    canonical: list[str | None] = [
        value if value in CANONICAL_COLUMNS else None for value in headers
    ]
    present = [value for value in canonical if value]
    if len(present) != len(set(present)):
        raise HTTPException(status_code=400, detail="duplicate_columns")
    if "email" not in present:
        raise HTTPException(status_code=400, detail="missing_email_column")

    rows: list[_RawRow] = []
    for index, values in enumerate(table[1:], start=2):
        mapped = {
            column: _as_text(values[position])
            for position, column in enumerate(canonical)
            if column and position < len(values)
        }
        if not any(mapped.values()):
            continue
        rows.append(_RawRow(row_number=index, values=mapped))
        if len(rows) > MAX_IMPORT_ROWS:
            raise HTTPException(status_code=400, detail="too_many_rows")
    if not rows:
        raise HTTPException(status_code=400, detail="empty_file")
    return rows


def _staff_emails(session: Session, gym_id: uuid.UUID) -> set[str]:
    users = session.exec(
        select(User)
        .join(GymMembership, col(GymMembership.user_id) == col(User.id))
        .join(
            GymMembershipRole,
            col(GymMembershipRole.membership_id) == col(GymMembership.id),
        )
        .where(
            GymMembership.gym_id == gym_id,
            col(GymMembershipRole.role).in_(["COACH", "GYM_ADMIN"]),
        )
    ).all()
    return {user.email.lower() for user in users}


def _existing_athlete_emails(session: Session, gym_id: uuid.UUID) -> set[str]:
    profiles = session.exec(
        select(AthleteProfile).where(AthleteProfile.gym_id == gym_id)
    ).all()
    user_ids = [profile.user_id for profile in profiles]
    if not user_ids:
        return set()
    users = session.exec(select(User).where(col(User.id).in_(user_ids))).all()
    return {user.email.lower() for user in users}


def _resolve_default_coach(
    session: Session,
    *,
    gym_id: uuid.UUID,
    coach_user_id: uuid.UUID | None,
) -> User | None:
    if not coach_user_id:
        return None
    coach = next(
        (item for item in coach_members(session, gym_id) if item.user_id == coach_user_id),
        None,
    )
    if not coach:
        raise HTTPException(status_code=400, detail="invalid_coach")
    return session.get(User, coach.user_id)


def _prepare_rows(
    session: Session,
    *,
    gym_id: uuid.UUID,
    filename: str,
    content: bytes,
    default_coach_user_id: uuid.UUID | None,
) -> list[_PreparedRow]:
    raw_rows = _parse_rows(filename, content)
    default_coach = _resolve_default_coach(
        session, gym_id=gym_id, coach_user_id=default_coach_user_id
    )
    coaches = {coach.email.lower(): coach for coach in coach_members(session, gym_id)}
    existing = _existing_athlete_emails(session, gym_id)
    staff = _staff_emails(session, gym_id)
    seen: set[str] = set()
    prepared: list[_PreparedRow] = []

    for raw in raw_rows:
        values = raw.values
        email = values.get("email", "").strip().lower()
        full_name = values.get("full_name") or None
        phone = values.get("phone") or None
        coach_email_value = (values.get("coach_email") or "").lower() or None
        training_goal = values.get("training_goal") or None
        messages: list[str] = []
        status: Literal["valid", "skip", "error"] = "valid"

        if not email:
            status = "error"
            messages.append("missing_email")
        else:
            try:
                email = str(_EMAIL.validate_python(email)).lower()
            except ValidationError:
                status = "error"
                messages.append("invalid_email")

        coach: User | None = None
        if coach_email_value:
            coach_public = coaches.get(coach_email_value)
            coach = (
                session.get(User, coach_public.user_id) if coach_public else None
            )
            if not coach:
                status = "error"
                messages.append("coach_not_found")
        else:
            coach = default_coach
            if not coach:
                status = "error"
                messages.append("coach_required")

        if full_name and len(full_name) > 255:
            status = "error"
            messages.append("full_name_too_long")
        if phone and len(phone) > 50:
            status = "error"
            messages.append("phone_too_long")
        if training_goal and len(training_goal) > 1000:
            status = "error"
            messages.append("goal_too_long")

        height_cm: float | None = None
        weight_kg: float | None = None
        training_days_count: int | None = None
        try:
            height_cm = _parse_number(values.get("height_cm", ""))
            if height_cm is not None and height_cm <= 0:
                raise ValueError
        except ValueError:
            status = "error"
            messages.append("invalid_height")
        try:
            weight_kg = _parse_number(values.get("weight_kg", ""))
            if weight_kg is not None and weight_kg <= 0:
                raise ValueError
        except ValueError:
            status = "error"
            messages.append("invalid_weight")
        days_value = values.get("training_days_count", "").translate(DIGITS).strip()
        if days_value:
            try:
                training_days_count = int(days_value)
                if not 1 <= training_days_count <= 7:
                    raise ValueError
            except ValueError:
                status = "error"
                messages.append("invalid_training_days")

        if email and status != "error" and email in staff:
            status = "error"
            messages.append("staff_email_conflict")
        elif email and status != "error" and email in seen:
            status = "skip"
            messages.append("duplicate_in_file")
        elif email and status != "error" and email in existing:
            status = "skip"
            messages.append("existing_athlete")
        if email:
            seen.add(email)

        prepared.append(
            _PreparedRow(
                row_number=raw.row_number,
                email=email,
                full_name=full_name,
                phone=phone,
                coach_email=coach_email_value,
                coach=coach,
                training_goal=training_goal,
                height_cm=height_cm,
                weight_kg=weight_kg,
                training_days_count=training_days_count,
                status=status,
                messages=messages,
            )
        )
    return prepared


def preview_athlete_import(
    session: Session,
    *,
    gym_id: uuid.UUID,
    filename: str,
    content: bytes,
    default_coach_user_id: uuid.UUID | None,
) -> AthleteImportPreview:
    rows = _prepare_rows(
        session,
        gym_id=gym_id,
        filename=filename,
        content=content,
        default_coach_user_id=default_coach_user_id,
    )
    coach_names = {
        item.user_id: item.full_name or item.email
        for item in coach_members(session, gym_id)
    }
    preview_rows = [
        AthleteImportPreviewRow(
            row_number=row.row_number,
            email=row.email,
            full_name=row.full_name,
            coach_email=row.coach.email if row.coach else row.coach_email,
            coach_name=coach_names.get(row.coach.id) if row.coach else None,
            status=row.status,
            messages=row.messages,
        )
        for row in rows
    ]
    return AthleteImportPreview(
        rows=preview_rows,
        total=len(preview_rows),
        valid=sum(1 for row in preview_rows if row.status == "valid"),
        skipped=sum(1 for row in preview_rows if row.status == "skip"),
        errors=sum(1 for row in preview_rows if row.status == "error"),
    )


def _create_athlete_from_row(
    session: Session,
    *,
    gym: Gym,
    current_user: User,
    row: _PreparedRow,
    activation_method: ActivationMethod,
    initial_password: str | None,
) -> None:
    assert row.coach is not None
    activation = activate_user(
        session,
        email=row.email,
        full_name=row.full_name,
        gym=gym,
        method=activation_method,
        initial_password=initial_password,
    )
    membership = ensure_membership(
        session, gym_id=gym.id, user_id=activation.user.id
    )
    session.flush()
    add_roles(session, membership=membership, roles=["ATHLETE"])
    if not activation.user.full_name and row.full_name:
        activation.user.full_name = row.full_name
        session.add(activation.user)
    profile = AthleteProfile(
        gym_id=gym.id,
        user_id=activation.user.id,
        phone=row.phone,
        height_cm=row.height_cm,
        weight_kg=row.weight_kg,
        training_goal=row.training_goal,
        training_days_count=row.training_days_count,
        training_days=[],
    )
    session.add(profile)
    session.flush()
    session.add(
        CoachAssignment(
            gym_id=gym.id,
            athlete_user_id=activation.user.id,
            coach_user_id=row.coach.id,
            assigned_by_id=current_user.id,
        )
    )
    session.flush()


def import_athletes(
    session: Session,
    *,
    gym: Gym,
    current_user: User,
    filename: str,
    content: bytes,
    default_coach_user_id: uuid.UUID | None,
    activation_method: ActivationMethod,
    initial_password: str | None,
) -> AthleteImportResult:
    if activation_method == "password" and (
        not initial_password or len(initial_password) < 8
    ):
        raise HTTPException(status_code=400, detail="initial_password_required")
    rows = _prepare_rows(
        session,
        gym_id=gym.id,
        filename=filename,
        content=content,
        default_coach_user_id=default_coach_user_id,
    )
    created = 0
    skipped = sum(1 for row in rows if row.status == "skip")
    failed: list[AthleteImportFailure] = []
    for row in rows:
        if row.status == "skip":
            continue
        if row.status == "error":
            failed.append(
                AthleteImportFailure(
                    row_number=row.row_number,
                    email=row.email,
                    messages=row.messages,
                )
            )
            continue
        try:
            with session.begin_nested():
                _create_athlete_from_row(
                    session,
                    gym=gym,
                    current_user=current_user,
                    row=row,
                    activation_method=activation_method,
                    initial_password=initial_password,
                )
            created += 1
        except HTTPException as exc:
            failed.append(
                AthleteImportFailure(
                    row_number=row.row_number,
                    email=row.email,
                    messages=[str(exc.detail)],
                )
            )
        except Exception:
            failed.append(
                AthleteImportFailure(
                    row_number=row.row_number,
                    email=row.email,
                    messages=["import_failed"],
                )
            )
    session.commit()
    return AthleteImportResult(
        total=len(rows),
        created=created,
        skipped=skipped,
        failed=failed,
    )


def _style_header(worksheet: Any, columns: list[str]) -> None:
    fill = PatternFill("solid", fgColor="DCEFE7")
    font = Font(bold=True)
    for index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=index, value=column)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
        worksheet.column_dimensions[get_column_letter(index)].width = max(
            16, len(column) + 6
        )
    worksheet.freeze_panes = "A2"


def build_template_csv() -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=CANONICAL_COLUMNS)
    writer.writeheader()
    writer.writerow(SAMPLE_ROW)
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def build_template_xlsx(session: Session, gym_id: uuid.UUID) -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Athletes"
    sheet.append(CANONICAL_COLUMNS)
    sheet.append([SAMPLE_ROW[column] for column in CANONICAL_COLUMNS])
    _style_header(sheet, CANONICAL_COLUMNS)
    sheet.row_dimensions[1].height = 24
    for row in sheet.iter_rows(min_row=2, max_row=MAX_IMPORT_ROWS + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    days_validation = DataValidation(
        type="whole", operator="between", formula1="1", formula2="7"
    )
    days_validation.error = "تعداد روز تمرین باید بین ۱ تا ۷ باشد."
    sheet.add_data_validation(days_validation)
    days_validation.add(f"H2:H{MAX_IMPORT_ROWS + 1}")

    coaches = coach_members(session, gym_id)
    coaches_sheet = workbook.create_sheet("Coaches")
    coaches_sheet.append(["coach_email", "coach_name"])
    for coach in coaches:
        coaches_sheet.append([coach.email, coach.full_name or ""])
    _style_header(coaches_sheet, ["coach_email", "coach_name"])
    coaches_sheet.column_dimensions["A"].width = 32
    coaches_sheet.column_dimensions["B"].width = 24
    if coaches:
        coach_validation = DataValidation(
            type="list",
            formula1=f"=Coaches!$A$2:$A${len(coaches) + 1}",
            allow_blank=True,
        )
        coach_validation.error = "ایمیل مربی باید در فهرست مربیان باشگاه باشد."
        sheet.add_data_validation(coach_validation)
        coach_validation.add(f"D2:D{MAX_IMPORT_ROWS + 1}")

    guide = workbook.create_sheet("راهنما")
    guide.append(["ستون", "اجباری", "توضیح", "نمونه"])
    for column, required, description, example in COLUMN_GUIDE:
        guide.append([column, "بله" if required else "خیر", description, example])
    _style_header(guide, ["ستون", "اجباری", "توضیح", "نمونه"])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 60
    guide.column_dimensions["D"].width = 28
    guide.append([])
    notes = [
        "ردیف نمونه را حذف کنید یا با داده‌های واقعی جایگزین کنید.",
        "ایمیل هر ورزشکار باید یکتا باشد؛ ردیف‌های تکراری یا موجود رد می‌شوند.",
        "اگر coach_email خالی باشد، مربی پیش‌فرض انتخاب‌شده در برنامه اعمال می‌شود.",
        f"حداکثر {MAX_IMPORT_ROWS} ردیف و {MAX_IMPORT_BYTES // (1024 * 1024)} مگابایت برای هر فایل.",
        "روش فعال‌سازی (دعوت یا رمز عبور) هنگام بارگذاری فایل انتخاب می‌شود.",
    ]
    for note in notes:
        guide.append([note])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
