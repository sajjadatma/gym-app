from __future__ import annotations

import io
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException, UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.routes.athletes import _read_import_upload
from app.services.athlete_import import MAX_IMPORT_BYTES
from tests.utils.helpers import (
    add_member,
    auth_headers,
    change_password,
    create_athlete,
    create_gym,
    gym_headers,
    unique_email,
)


def _setup(client: TestClient, superuser_headers: dict[str, str]):
    owner_email = unique_email("owner")
    gym = create_gym(
        client,
        superuser_headers,
        name="Import Gym",
        owner_email=owner_email,
        password="owner-pass-123",
    )["gym"]
    raw = auth_headers(client, owner_email, "owner-pass-123")
    change_password(
        client,
        raw,
        current_password="owner-pass-123",
        new_password="owner-pass-456",
    )
    owner = gym_headers(
        client, email=owner_email, password="owner-pass-456", gym_id=gym["id"]
    )
    coach_email = unique_email("coach")
    coach_member = add_member(
        client,
        owner,
        gym["id"],
        email=coach_email,
        roles=["COACH"],
        password="coach-pass-123",
    )["member"]
    coach_raw = auth_headers(client, coach_email, "coach-pass-123")
    change_password(
        client,
        coach_raw,
        current_password="coach-pass-123",
        new_password="coach-pass-456",
    )
    coach = gym_headers(
        client, email=coach_email, password="coach-pass-456", gym_id=gym["id"]
    )
    return gym, owner, coach, coach_member, coach_email


def _csv(*rows: list[str]) -> bytes:
    header = (
        "email,full_name,phone,coach_email,training_goal,height_cm,"
        "weight_kg,training_days_count"
    )
    return ("\ufeff" + "\n".join([header, *[",".join(row) for row in rows]])).encode(
        "utf-8"
    )


def test_import_templates_are_available_to_managers_only(
    client: TestClient, superuser_headers: dict[str, str]
) -> None:
    _gym, owner, coach, _coach_member, _coach_email = _setup(
        client, superuser_headers
    )

    csv_response = client.get(
        "/api/v1/athletes/import/template", headers=owner
    )
    assert csv_response.status_code == 200, csv_response.text
    assert "text/csv" in csv_response.headers["content-type"]
    assert "attachment" in csv_response.headers["content-disposition"]
    assert csv_response.content.startswith("\ufeffemail,full_name".encode())
    assert "athlete@example.com" in csv_response.content.decode("utf-8-sig")

    xlsx_response = client.get(
        "/api/v1/athletes/import/template?file_format=xlsx", headers=owner
    )
    assert xlsx_response.status_code == 200, xlsx_response.text
    workbook = load_workbook(io.BytesIO(xlsx_response.content))
    assert {"Athletes", "Coaches", "راهنما"} <= set(workbook.sheetnames)
    athletes = workbook["Athletes"]
    assert athletes["A1"].value == "email"
    assert athletes["A2"].value == "athlete@example.com"

    forbidden = client.get("/api/v1/athletes/import/template", headers=coach)
    assert forbidden.status_code == 403


def test_preview_and_import_report_valid_skipped_and_error_rows(
    client: TestClient, superuser_headers: dict[str, str]
) -> None:
    gym, owner, coach, coach_member, coach_email = _setup(
        client, superuser_headers
    )
    existing_email = unique_email("existing")
    create_athlete(
        client,
        coach,
        email=existing_email,
        coach_user_id=coach_member["user_id"],
        password="athlete-pass-123",
    )

    valid_one = unique_email("bulk-one")
    valid_two = unique_email("bulk-two")
    duplicate = unique_email("bulk-duplicate")
    invalid = "not-an-email"
    missing_coach = unique_email("bulk-coach")
    content = _csv(
        [valid_one, "Bulk One", "09120000001", coach_email, "قدرت", "180", "80", "3"],
        [valid_two, "Bulk Two", "09120000002", "", "حجم", "181", "81", "2"],
        [duplicate, "Duplicate", "", coach_email, "", "", "", ""],
        [duplicate, "Duplicate again", "", coach_email, "", "", "", ""],
        [invalid, "Invalid", "", coach_email, "", "", "", ""],
        [existing_email, "Existing", "", coach_email, "", "", "", ""],
        [missing_coach, "No coach", "", "missing@example.com", "", "", "", ""],
    )
    files = {"file": ("athletes.csv", content, "text/csv")}
    preview = client.post(
        "/api/v1/athletes/import/preview",
        headers=owner,
        files=files,
        data={"default_coach_user_id": coach_member["user_id"]},
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["total"] == 7
    assert body["valid"] == 3
    assert body["skipped"] == 2
    assert body["errors"] == 2
    statuses = [row["status"] for row in body["rows"]]
    assert statuses == ["valid", "valid", "valid", "skip", "error", "skip", "error"]
    assert body["rows"][3]["messages"] == ["duplicate_in_file"]
    assert body["rows"][4]["messages"] == ["invalid_email"]
    assert body["rows"][6]["messages"] == ["coach_not_found"]
    assert body["rows"][1]["coach_email"] == coach_email

    imported = client.post(
        "/api/v1/athletes/import",
        headers=owner,
        files=files,
        data={
            "default_coach_user_id": coach_member["user_id"],
            "activation_method": "password",
            "initial_password": "BulkPass123!",
        },
    )
    assert imported.status_code == 200, imported.text
    result = imported.json()
    assert result["created"] == 3
    assert result["skipped"] == 2
    assert len(result["failed"]) == 2

    athletes = client.get("/api/v1/athletes", headers=owner).json()["data"]
    emails = {row["email"] for row in athletes}
    assert {valid_one, valid_two, duplicate, existing_email} <= emails
    assert invalid not in emails
    imported_profile = next(row for row in athletes if row["email"] == valid_one)
    assert imported_profile["coach_user_id"] == coach_member["user_id"]


def test_xlsx_upload_preview(client: TestClient, superuser_headers: dict[str, str]) -> None:
    _gym, owner, _coach, coach_member, coach_email = _setup(
        client, superuser_headers
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "email",
            "full_name",
            "phone",
            "coach_email",
            "training_goal",
            "height_cm",
            "weight_kg",
            "training_days_count",
        ]
    )
    email = unique_email("xlsx")
    sheet.append([email, "XLSX Athlete", "", coach_email, "", "175", "70", "4"])
    output = io.BytesIO()
    workbook.save(output)

    response = client.post(
        "/api/v1/athletes/import/preview",
        headers=owner,
        files={
            "file": (
                "athletes.xlsx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"default_coach_user_id": coach_member["user_id"]},
    )
    assert response.status_code == 200, response.text
    assert response.json()["valid"] == 1
    assert response.json()["rows"][0]["email"] == email


def test_import_rejects_staff_emails(
    client: TestClient, superuser_headers: dict[str, str]
) -> None:
    _gym, owner, _coach, coach_member, coach_email = _setup(
        client, superuser_headers
    )
    content = _csv([coach_email, "Coach as athlete", "", coach_email, "", "", "", ""])
    response = client.post(
        "/api/v1/athletes/import/preview",
        headers=owner,
        files={"file": ("athletes.csv", content, "text/csv")},
        data={"default_coach_user_id": coach_member["user_id"]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["errors"] == 1
    assert body["rows"][0]["messages"] == ["staff_email_conflict"]


def test_import_upload_read_is_bounded() -> None:
    upload = UploadFile(filename="athletes.csv", file=io.BytesIO())
    content = b"x" * MAX_IMPORT_BYTES
    upload.read = AsyncMock(return_value=content)  # type: ignore[method-assign]
    assert pytest.importorskip("asyncio").run(_read_import_upload(upload)) == content
    upload.read.assert_awaited_once_with(MAX_IMPORT_BYTES + 1)

    oversized = UploadFile(filename="athletes.csv", file=io.BytesIO())
    oversized.read = AsyncMock(return_value=b"x" * (MAX_IMPORT_BYTES + 1))  # type: ignore[method-assign]
    with pytest.raises(HTTPException) as exc:
        pytest.importorskip("asyncio").run(_read_import_upload(oversized))
    assert exc.value.status_code == 400
    assert exc.value.detail == "import_file_too_large"


@pytest.mark.parametrize("path", ["/api/v1/athletes/import/preview", "/api/v1/athletes/import"])
def test_preview_and_import_reject_oversized_uploads(
    client: TestClient, superuser_headers: dict[str, str], path: str
) -> None:
    _gym, owner, _coach, coach_member, _coach_email = _setup(
        client, superuser_headers
    )
    response = client.post(
        path,
        headers=owner,
        files={"file": ("athletes.csv", b"x" * (MAX_IMPORT_BYTES + 1), "text/csv")},
        data={
            "default_coach_user_id": coach_member["user_id"],
            "activation_method": "password",
            "initial_password": "BulkPass123!",
        },
    )
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == "import_file_too_large"


@pytest.mark.parametrize("path", ["/api/v1/athletes/import/preview", "/api/v1/athletes/import"])
def test_preview_and_import_report_malformed_csv(
    client: TestClient, superuser_headers: dict[str, str], path: str
) -> None:
    _gym, owner, _coach, coach_member, _coach_email = _setup(
        client, superuser_headers
    )
    response = client.post(
        path,
        headers=owner,
        files={"file": ("athletes.csv", b'email,full_name\n"unterminated', "text/csv")},
        data={
            "default_coach_user_id": coach_member["user_id"],
            "activation_method": "password",
            "initial_password": "BulkPass123!",
        },
    )
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == "invalid_file"


@pytest.mark.parametrize("path", ["/api/v1/athletes/import/preview", "/api/v1/athletes/import"])
@pytest.mark.parametrize("content", [b"", _csv()])
def test_preview_and_import_reject_empty_files(
    client: TestClient,
    superuser_headers: dict[str, str],
    path: str,
    content: bytes,
) -> None:
    _gym, owner, _coach, coach_member, _coach_email = _setup(
        client, superuser_headers
    )
    response = client.post(
        path,
        headers=owner,
        files={"file": ("athletes.csv", content, "text/csv")},
        data={
            "default_coach_user_id": coach_member["user_id"],
            "activation_method": "password",
            "initial_password": "BulkPass123!",
        },
    )
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == "empty_file"


@pytest.mark.parametrize("path", ["/api/v1/athletes/import/preview", "/api/v1/athletes/import"])
def test_import_preview_and_write_require_gym_admin(
    client: TestClient, superuser_headers: dict[str, str], path: str
) -> None:
    _gym, _owner, coach, coach_member, _coach_email = _setup(
        client, superuser_headers
    )
    email = unique_email("forbidden-import")
    response = client.post(
        path,
        headers=coach,
        files={
            "file": (
                "athletes.csv",
                _csv([email, "Forbidden", "", "", "", "", "", ""]),
                "text/csv",
            )
        },
        data={
            "default_coach_user_id": coach_member["user_id"],
            "activation_method": "password",
            "initial_password": "BulkPass123!",
        },
    )
    assert response.status_code == 403, response.text
    assert response.json()["detail"] == "gym_admin_required"


@pytest.mark.parametrize("path", ["/api/v1/athletes/import/preview", "/api/v1/athletes/import"])
def test_import_rejects_foreign_gym_default_coach(
    client: TestClient, superuser_headers: dict[str, str], path: str
) -> None:
    _gym_a, owner_a, _coach_a, _member_a, _email_a = _setup(
        client, superuser_headers
    )
    _gym_b, _owner_b, _coach_b, coach_member_b, _email_b = _setup(
        client, superuser_headers
    )
    response = client.post(
        path,
        headers=owner_a,
        files={
            "file": (
                "athletes.csv",
                _csv([unique_email("foreign-coach"), "Foreign", "", "", "", "", "", ""]),
                "text/csv",
            )
        },
        data={
            "default_coach_user_id": coach_member_b["user_id"],
            "activation_method": "password",
            "initial_password": "BulkPass123!",
        },
    )
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == "invalid_coach"
